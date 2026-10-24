"""
Report management: auto-save, plain-text save, and Word (.docx) export.
Word export requires:  pip install python-docx
"""

from __future__ import annotations

import io
import logging
import os
import re
from datetime import datetime
from typing import TYPE_CHECKING, Optional

from src.core.patient_schema import PATIENT_FIELDS
from src.core.paths import autosave_dir, report_filename

if TYPE_CHECKING:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

logger = logging.getLogger(__name__)

try:
    from docx import Document  # type: ignore[assignment]
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[assignment]
    from docx.shared import Inches, Pt, RGBColor  # type: ignore[assignment]

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning(
        "python-docx not installed – Word export unavailable. "
        "Run: pip install python-docx"
    )

# Section names a report line can be (with or without a trailing colon).
_SECTION_HEADERS = {
    "TECHNIQUE", "FINDINGS", "IMPRESSION", "CONCLUSION", "COMPARISON",
    "CLINICAL INDICATION", "CLINICAL DETAILS", "EXAMINATION",           # RSNA
    "CLINICAL HISTORY", "OPINION",                                      # RCR
    "CLINICAL INFORMATION", "RECOMMENDATION", "RECOMMENDATIONS",        # ECR/ESR
}
_LEADING_HEADER = re.compile(
    rf"^({'|'.join(sorted(_SECTION_HEADERS, key=len, reverse=True))})\s*:\s*(.*)$", re.IGNORECASE
)
#: Characters XML 1.0 cannot hold (every control character but tab and line
#: breaks). One stray one, a barcode scanner's ESC in the ID box, made the
#: Word export raise while the text export of the same report worked.
_NOT_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _clean(value: object) -> str:
    """Every string either export writes, report or header, goes through here."""
    return _NOT_XML.sub("", str(value or ""))


# ---------------------------------------------------------------------------
# Auto-save
# ---------------------------------------------------------------------------

def autosave_report(text: str, patient_info: dict) -> Optional[str]:
    """Write report to autosave folder.  Returns saved path, or None on error."""
    if not text.strip():
        return None
    try:
        filename = report_filename(
            patient_info, ".txt", prefix="dictation_", fallback_id="unknown"
        )
        path = os.path.join(str(autosave_dir()), filename)
        with open(path, "w", encoding="utf-8") as f:
            f.write(format_plain_text_report(text, patient_info))
        return path
    except Exception as exc:
        logger.warning("Auto-save failed: %s", exc)
        return None


# ---------------------------------------------------------------------------
# Plain-text save
# ---------------------------------------------------------------------------

def save_report_txt(path: str, text: str, patient_info: dict) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.write(format_plain_text_report(text, patient_info))


def format_plain_text_report(text: str, patient_info: dict) -> str:
    """Return the formatted plain-text report body."""
    text = _clean(text)
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    rule = "=" * 64
    label_width = max(len(label) for _, label in PATIENT_FIELDS) + 2
    header = [
        "RADIOLOGY REPORT",
        rule,
        *(
            f"{label + ':':<{label_width}}{_clean(patient_info.get(key))}"
            for key, label in PATIENT_FIELDS
        ),
        rule,
        "",
    ]
    footer = [
        "",
        rule,
        f"Reported: {now}",
        "Reporting Radiologist: _________________________________",
    ]
    return "\n".join(header) + "\n" + text + "\n" + "\n".join(footer)



# ---------------------------------------------------------------------------
# Word (.docx) export
# ---------------------------------------------------------------------------

def _build_word_document(text: str, patient_info: dict):
    doc = Document()

    # ---- Page margins ----
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # ---- Title ----
    title = doc.add_heading("RADIOLOGY REPORT", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- Patient info table (two label/value pairs per row) ----
    rows = (len(PATIENT_FIELDS) + 1) // 2
    tbl = doc.add_table(rows=rows, cols=4)
    tbl.style = "Table Grid"
    for index, (key, label) in enumerate(PATIENT_FIELDS):
        row, col = divmod(index, 2)
        _tbl_cell(tbl, row, col * 2, f"{label}:", bold=True)
        _tbl_cell(tbl, row, col * 2 + 1, _clean(patient_info.get(key)))

    doc.add_paragraph()  # spacer

    # ---- Report body ----
    # A heading is a whole line that names a section. "FINDINGS: normal" is a
    # paragraph led by its label; "Findings are within normal limits" is just
    # a paragraph (matching on how a line began turned it into a heading).
    for line in _clean(text).split("\n"):
        line = line.rstrip()
        if not line.strip():
            doc.add_paragraph()
            continue
        name = line.strip().rstrip(":").strip().upper()
        if name in _SECTION_HEADERS or (
            line.strip().endswith(":") and line.strip().upper() == line.strip()
            and len(line.split()) <= 5
        ):
            h = doc.add_heading(line.strip(), level=2)
            h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            continue
        p = doc.add_paragraph()
        led = _LEADING_HEADER.match(line.strip())
        if led:
            p.add_run(f"{led.group(1)}: ").bold = True
            p.add_run(led.group(2))
        else:
            p.add_run(line)  # leading indentation kept
        for run in p.runs:
            run.font.size = Pt(11)

    # ---- Signature line ----
    doc.add_paragraph()
    now = datetime.now().strftime("%d/%m/%Y")
    sig = doc.add_paragraph()
    sig.add_run("Reporting Radiologist: ").bold = True
    sig.add_run("_________________________________")
    sig.add_run(f"\t\t\tDate: {now}")

    return doc


def export_to_word_bytes(text: str, patient_info: dict) -> bytes:
    """Render the Word report into an in-memory ``.docx`` payload."""
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "python-docx is not installed.\n"
            "Install it with:  pip install python-docx"
        )

    doc = _build_word_document(text, patient_info)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_to_word(path: str, text: str, patient_info: dict) -> None:
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "python-docx is not installed.\n"
            "Install it with:  pip install python-docx"
        )

    doc = _build_word_document(text, patient_info)
    doc.save(path)


def _tbl_cell(table, row: int, col: int, text: str, bold: bool = False) -> None:
    cell = table.cell(row, col)
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
