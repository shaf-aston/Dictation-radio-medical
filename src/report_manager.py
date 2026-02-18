"""
Report management: auto-save, plain-text save, and Word (.docx) export.
Word export requires:  pip install python-docx
"""
import os
import logging
from datetime import datetime
from typing import Optional

logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not installed – Word export unavailable. Run: pip install python-docx")


# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------

def get_autosave_dir() -> str:
    base = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    path = os.path.join(base, "autosave")
    os.makedirs(path, exist_ok=True)
    return path


# ---------------------------------------------------------------------------
# Auto-save
# ---------------------------------------------------------------------------

def autosave_report(text: str, patient_info: dict) -> Optional[str]:
    """Write report to autosave folder.  Returns saved path, or None on error."""
    if not text.strip():
        return None
    try:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        pid = (patient_info.get("id") or "unknown").strip().replace(" ", "_") or "unknown"
        filename = f"dictation_{pid}_{ts}.txt"
        path = os.path.join(get_autosave_dir(), filename)
        with open(path, "w", encoding="utf-8") as f:
            f.write(_format_plain_text(text, patient_info))
        return path
    except Exception as exc:
        logger.warning("Auto-save failed: %s", exc)
        return None


# ---------------------------------------------------------------------------
# Plain-text save
# ---------------------------------------------------------------------------

def save_report_txt(path: str, text: str, patient_info: dict) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.write(_format_plain_text(text, patient_info))


def _format_plain_text(text: str, patient_info: dict) -> str:
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    header = [
        "RADIOLOGY REPORT",
        "=" * 64,
        f"Patient:      {patient_info.get('name', '')}",
        f"Patient ID:   {patient_info.get('id', '')}",
        f"Date of Birth:{patient_info.get('dob', '')}",
        f"Study Date:   {patient_info.get('study_date', '')}",
        f"Referring:    {patient_info.get('referring', '')}",
        f"Accession #:  {patient_info.get('accession', '')}",
        "=" * 64,
        "",
    ]
    footer = [
        "",
        "=" * 64,
        f"Reported: {now}",
        "Reporting Radiologist: _________________________________",
    ]
    return "\n".join(header) + text + "\n".join(footer)


# ---------------------------------------------------------------------------
# Word (.docx) export
# ---------------------------------------------------------------------------

def export_to_word(path: str, text: str, patient_info: dict) -> None:
    if not DOCX_AVAILABLE:
        raise RuntimeError(
            "python-docx is not installed.\n"
            "Install it with:  pip install python-docx"
        )

    doc = Document()

    # ---- Page margins ----
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # ---- Title ----
    title = doc.add_heading("RADIOLOGY REPORT", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- Patient info table ----
    tbl = doc.add_table(rows=3, cols=4)
    tbl.style = "Table Grid"
    _tbl_cell(tbl, 0, 0, "Patient Name:", bold=True)
    _tbl_cell(tbl, 0, 1, patient_info.get("name", ""))
    _tbl_cell(tbl, 0, 2, "Patient ID:", bold=True)
    _tbl_cell(tbl, 0, 3, patient_info.get("id", ""))
    _tbl_cell(tbl, 1, 0, "Date of Birth:", bold=True)
    _tbl_cell(tbl, 1, 1, patient_info.get("dob", ""))
    _tbl_cell(tbl, 1, 2, "Study Date:", bold=True)
    _tbl_cell(tbl, 1, 3, patient_info.get("study_date", ""))
    _tbl_cell(tbl, 2, 0, "Referring Clinician:", bold=True)
    _tbl_cell(tbl, 2, 1, patient_info.get("referring", ""))
    _tbl_cell(tbl, 2, 2, "Accession #:", bold=True)
    _tbl_cell(tbl, 2, 3, patient_info.get("accession", ""))

    doc.add_paragraph()  # spacer

    # ---- Report body ----
    _SECTION_HEADERS = {
        "TECHNIQUE", "FINDINGS", "IMPRESSION", "CONCLUSION",
        "CLINICAL INDICATION", "CLINICAL DETAILS", "COMPARISON",
        "TECHNIQUE:", "FINDINGS:", "IMPRESSION:", "CONCLUSION:",
        "CLINICAL INDICATION:", "CLINICAL DETAILS:", "COMPARISON:",
    }

    for line in text.split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue

        upper = stripped.upper()
        # Detect section headers: all-caps words or known headers ending with ':'
        is_header = (
            upper in _SECTION_HEADERS
            or (stripped.endswith(":") and stripped.upper() == stripped and len(stripped.split()) <= 4)
            or any(upper.startswith(h) for h in _SECTION_HEADERS)
        )

        if is_header:
            h = doc.add_heading(stripped, level=2)
            h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
        else:
            p = doc.add_paragraph(stripped)
            for run in p.runs:
                run.font.size = Pt(11)

    # ---- Signature line ----
    doc.add_paragraph()
    now = datetime.now().strftime("%d/%m/%Y")
    sig = doc.add_paragraph()
    sig.add_run("Reporting Radiologist: ").bold = True
    sig.add_run("_________________________________")
    sig.add_run(f"\t\t\tDate: {now}")

    doc.save(path)


def _tbl_cell(table, row: int, col: int, text: str, bold: bool = False) -> None:
    cell = table.cell(row, col)
    para = cell.paragraphs[0]
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(10)
